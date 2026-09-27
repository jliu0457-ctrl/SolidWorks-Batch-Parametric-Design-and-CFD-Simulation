#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Run-FlowSample 里**不需要 SolidWorks** 的部分：坐标变换、状态机、续跑守卫、目标解析。

坐标变换那几条特别值得测 —— 布局猜错了整条空间识别都会静默跑偏，
而唯一的症状是「射线打到了别的零件」，排查成本极高。
这里的期望值全部来自 _analysis/co_11蝶板.json 的真实变换矩阵。
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

MAPPING = Path(__file__).resolve().parents[1]
#: 根目录的人工工程（`流体仿真2/1/`）—— 只在需要拿真实数据当靶子的测试里用。
HUMAN_DIR = MAPPING.parent / "1"
sys.path.insert(0, str(MAPPING / "scripts"))

import flow_geometry as fg  # noqa: E402


def _load_orchestrator():
    """脚本名带连字符（`Run-FlowSample.py`），不是合法标识符，只能用 importlib 装。"""
    path = MAPPING / "scripts" / "Run-FlowSample.py"
    spec = importlib.util.spec_from_file_location("Run_FlowSample", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules["Run_FlowSample"] = module
    spec.loader.exec_module(module)
    return module


R = _load_orchestrator()

# 实测的组件变换（_analysis/co_11蝶板.json:final_transforms）
T_INLET_PIPE = [0, 0, 1, 0, 1, 0, -1, 0, 0, -0.115, 0, 0, 1, 0, 0, 0]
T_BODY = [1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0]
T_DISC_45 = [0.70710678, -0.70710678, 0, 0.70710678, 0.70710678, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0]
T_SHAFT = [0.0, 0.0, -1, -0.70710678, -0.70710678, -0.0, -0.70710678, 0.70710678, 0.0,
           0.002616, 0.002616, 0.5385, 1, 0, 0, 0]


# ---------------------------------------------------------------- 坐标变换

def test_identity_transform_leaves_points_alone():
    assert R.transform_point(T_BODY, [1.0, 2.0, 3.0]) == pytest.approx([1.0, 2.0, 3.0])


def test_inlet_pipe_local_z_maps_to_world_minus_x():
    """入口管道：局部 Z 轴映射到世界 -X，沿 +Z 伸出 0.5 m → 世界 x=-0.615。

    这正是 open_circular_edges.json 里管口圆边的位置 —— 变换写错了就对不上。
    """
    assert R.transform_vector(T_INLET_PIPE, [0, 0, 1]) == pytest.approx([-1.0, 0.0, 0.0], abs=1e-9)
    assert R.transform_vector(T_INLET_PIPE, [0, 1, 0]) == pytest.approx([0.0, 1.0, 0.0], abs=1e-9)
    # 原点在管子内端 x=-0.115；沿局部 +Z 走 0.5 m 到外端
    assert R.transform_point(T_INLET_PIPE, [0, 0, 0]) == pytest.approx([-0.115, 0.0, 0.0])
    assert R.transform_point(T_INLET_PIPE, [0, 0, 0.5]) == pytest.approx([-0.615, 0.0, 0.0])


def test_disc_transform_is_a_45_degree_rotation():
    """蝶板的组件变换就是「开度45°」那 45° 旋转 —— 转出来的角必须真是 45°。"""
    x_axis = R.transform_vector(T_DISC_45, [1, 0, 0])
    y_axis = R.transform_vector(T_DISC_45, [0, 1, 0])
    assert x_axis == pytest.approx([0.70710678, -0.70710678, 0], abs=1e-6)
    assert y_axis == pytest.approx([0.70710678, 0.70710678, 0], abs=1e-6)
    # 点积为 0 → 正交；世界原点不动（该几何下平移为 0）
    assert sum(a * b for a, b in zip(x_axis, y_axis)) == pytest.approx(0.0, abs=1e-9)
    assert R.transform_point(T_DISC_45, [0, 0, 0]) == pytest.approx([0.0, 0.0, 0.0])


def test_transform_vector_normalizes_but_point_does_not():
    scaled = [2, 0, 0, 0, 2, 0, 0, 0, 2, 1, 2, 3, 1, 0, 0, 0]
    assert R.transform_vector(scaled, [1, 0, 0]) == pytest.approx([1.0, 0.0, 0.0])
    assert R.transform_point(scaled, [1, 0, 0]) == pytest.approx([3.0, 2.0, 3.0])


def test_transformed_box_handles_a_rotation():
    """局部 ±0.1 的立方体经 45° 旋转后，外包围盒半宽变成 0.1*sqrt(2)。"""
    box = R.transformed_box(T_DISC_45, [-0.1, -0.1, -0.1, 0.1, 0.1, 0.1])
    assert box[0] == pytest.approx(-0.1 * 2 ** 0.5, abs=1e-6)
    assert box[3] == pytest.approx(0.1 * 2 ** 0.5, abs=1e-6)
    assert box[2] == pytest.approx(-0.1), "Z 方向不受绕 Z 旋转影响"


def test_model_bounds_combines_component_faces(monkeypatch):
    class Component:
        def __init__(self, box):
            self.box = box

    components = [Component([-0.6, -0.1, -0.2, -0.5, 0.1, 0.2]),
                  Component([1.0, -0.08, -0.15, 1.1, 0.08, 0.15])]
    identity = [1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0]
    monkeypatch.setattr(R.swa, "component_list", lambda _doc: components)
    monkeypatch.setattr(R.swa, "component_transform", lambda _comp: identity)
    monkeypatch.setattr(R.swa, "face_local_records", lambda comp: [{"box_m": comp.box}])

    assert R.model_bounds_m(object()) == pytest.approx([-0.6, -0.1, -0.2, 1.1, 0.1, 0.2])


def test_array_data_rejects_short_matrices():
    class Short:
        ArrayData = [1, 0, 0, 0, 1, 0]
    with pytest.raises(RuntimeError, match="不足 12 项"):
        R.array_data(Short())


def test_opening_edge_enumeration_skips_unrelated_components(monkeypatch):
    """S3 只需要两根管道和阀体，不应再跨 COM 扫描蝶板等零件的圆边。"""
    class Component:
        def __init__(self, name):
            self.Name2 = name

    components = [Component("入口管道^装配体-1"), Component("出口管道^装配体-1"),
                  Component("8英寸-03阀体-1"), Component("11蝶板-1")]
    scanned = []
    identity = [1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0, 1, 0, 0, 0]

    class FakeSwa:
        @staticmethod
        def component_list(_document):
            return components

        @staticmethod
        def safe_get(obj, name, default=None):
            return getattr(obj, name, default)

        @staticmethod
        def total_transform(_component):
            return identity

        @staticmethod
        def transform_array(transform):
            return transform

        @staticmethod
        def circular_edges(component):
            scanned.append(component.Name2)
            return [{"center_m": [0, 0, 0], "normal": [1, 0, 0],
                     "radius_m": 0.1, "closed": True}]

    monkeypatch.setattr(R, "swa", FakeSwa)

    rows = R.enumerate_open_circular_edges(object())

    assert len(rows) == 3
    assert scanned == [c.Name2 for c in components[:3]]
    assert all("11蝶板" not in row["component"] for row in rows)


# ---------------------------------------------------------------- 目标解析

def test_cap_goal_resolves_to_the_stable_entity_name():
    """挂封盖的目标解析出 face_name；组件名要给两种写法。"""
    binding = {"paired": {"inlet": {"component": '封盖3^8“D94R3Y-00组装图-1'}}}
    candidates, face_name = R.resolve_goal_components(
        {"name": "SG CV入口静压", "faces_on": ["FLOW_INLET_INNER"]}, binding, None)
    assert face_name == "FLOW_INLET_INNER"
    assert '封盖3^8“D94R3Y-00组装图-1' in candidates
    assert any("<" in c for c in candidates), "Flow 的引用串用的是 名字<序号> 形式"


def test_seat_goal_reports_no_candidates():
    """挂密封面的目标目前解析不出来 —— 必须如实返回空，让调用方报 needs_seat_faces。"""
    for spec in ({"name": "SG 密比压 平均", "faces_on": ["08大垫片", "11蝶板", "09密封圈"]},
                 {"name": "SG 力矩Z", "faces_on": ["04阀轴"], "components": ["11蝶板"]}):
        candidates, face_name = R.resolve_goal_components(spec, {"paired": {}}, None)
        assert candidates == [] and face_name is None, spec["name"]


def test_cap_name_candidates_give_both_spellings():
    """组件名有两种写法：装配体的 `封盖3-1` 和 Flow 引用串用的 `封盖3<1>`。"""
    got = R._cap_name_candidates("封盖3-1")
    assert got == list(dict.fromkeys(got)), "不能有重复"
    assert "封盖3-1" in got
    assert "封盖3<1>" in got


# ------------------------------------------------- 按面序号选面（人工那套面的正路）

def _rec(kind: str, area_mm2: float = 1000.0) -> dict:
    return {"area_m2": area_mm2 / 1e6, "is_plane": kind == "plane",
            "is_cone": kind == "cone", "is_cylinder": kind == "cylinder"}


class _FakeFaceObject:
    def __init__(self, label: str):
        self.label = label


class _FakeSwa:
    """`sw_api` 的替身：面记录按组件查，select_face 只记账。"""

    def __init__(self, faces: dict[str, list[dict]]):
        self.faces = faces
        self.selected: list[str] = []
        self.appends: list[bool] = []

    def face_local_records(self, comp):
        return self.faces[comp]

    def face_objects(self, comp):
        return [_FakeFaceObject(f"{comp}#{i}") for i in range(len(self.faces[comp]))]

    def select_face(self, face, append=False):
        self.appends.append(append)
        self.selected.append(face.label)
        return True


class _FakeNative:
    def __init__(self):
        self.cleared = 0

    def ClearSelection2(self, _flag):
        self.cleared += 1
        return True


@pytest.fixture()
def fake_native(monkeypatch):
    """把 `component_by_token` 和 `swa` 换成替身，返回 (native, swa)。"""
    faces = {
        "08大垫片": [_rec("plane", 11839.107), _rec("plane", 11696.414),
                     _rec("cylinder", 329.867), _rec("cone", 463.210)],
        "11蝶板": [_rec("other")] * 54 + [_rec("plane", 5376.589)],
    }
    fake = _FakeSwa(faces)
    monkeypatch.setattr(R, "swa", fake)
    monkeypatch.setattr(R, "component_by_token", lambda native, token: token)
    return _FakeNative(), fake


def test_face_index_selector_picks_the_configured_indices(fake_native):
    """大垫片 0/1 + 蝶板 54 —— 就是人工 `SG 密比压` 的那三个面（外加密封圈面 0）。"""
    native, fake = fake_native
    select = R.make_face_index_selector(native, [
        {"component": "08大垫片", "indices": [0, 1],
         "kinds": ["plane", "plane"], "baseline_area_mm2": [11839.107, 11696.414]},
        {"component": "11蝶板", "indices": [54],
         "kinds": ["plane"], "baseline_area_mm2": [5376.589]},
    ])
    assert select() == 3
    assert fake.selected == ["08大垫片#0", "08大垫片#1", "11蝶板#54"]
    assert fake.appends == [False, True, True], "第一个面用 append=False，其余 True"
    assert native.cleared == 1
    assert [d["kind"] for d in select.diagnostic] == ["plane", "plane", "plane"]
    assert select.diagnostic[0]["area_ratio"] == pytest.approx(1.0, abs=1e-4)
    assert select.last_picked == ["08大垫片#0", "08大垫片#1", "11蝶板#54"]


def test_face_index_selector_refuses_a_kind_mismatch(fake_native):
    """序号对了但面类型不对 = 零件拓扑变了。宁可炸，也不要静默绑错面。

    实测的代价：按半锥角挑密封面，三个面全挑错，密比压 841 网格面 vs 人工 288。
    """
    native, fake = fake_native
    select = R.make_face_index_selector(native, [
        {"component": "08大垫片", "indices": [3], "kinds": ["plane"]}])
    with pytest.raises(R.FaceSignatureMismatch, match="拓扑变了"):
        select()
    assert fake.selected == [], "类型不符时一个面都不该被选中"


def test_face_index_selector_refuses_an_out_of_range_index(fake_native):
    native, _ = fake_native
    select = R.make_face_index_selector(native, [
        {"component": "08大垫片", "indices": [99], "kinds": ["plane"]}])
    with pytest.raises(R.FaceSignatureMismatch, match="只有 4 个面"):
        select()


def test_face_index_selector_without_kinds_only_records_the_signature(fake_native):
    """`kinds` 是可选的 —— 缺了就不当判据，但仍要记下类型和面积比供审计。"""
    native, _ = fake_native
    select = R.make_face_index_selector(native, [
        {"component": "08大垫片", "indices": [2], "baseline_area_mm2": [329.867]}])
    assert select() == 1
    assert select.diagnostic[0]["kind"] == "cylinder"
    assert select.diagnostic[0]["area_ratio"] == pytest.approx(1.0, abs=1e-4)


# ------------------------------------------------- Flow 自己算的平均值（Goals.DAT）

#: 冻结的 `Goals.DAT` 样本 —— 值是 2026-09-21 从人工工程的界面「目标图」上**逐位抄下来**
#: 再和文件对照过的（数值/平均值/最小值/最大值/增量/标准 六列全对上）。
#: ⚠️ **测试必须用这个冻结样本，不能去读根目录那个活文件** ——
#: 用户会反复重跑，期望值会跟着变，测试就变成"追移动靶"（实测踩过三次）。
_GOALS_DAT_FIXTURE = """\
Iteration\tCPUTime\tPhysTime\tTravels\tValue\tAvValue\tMinValue\tMaxValue\tDelta\tCriteria\tPrevAvRefValue\tProgress\tCriteriaType\tCriteriaVarType\tCriteriaPercentage
1\t1\t0\t0.025\t326779.89\t326779.89\t326779.89\t326779.89\t0\t1633.89945\t0\t0\t1\t1\t0.5
160\t32\t0\t4\t213169.234\t213911.637\t213169.234\t214405.497\t447.64865\t1065.84617\t0\t100\t1\t1\t0.5
"""


def test_goals_dat_parser_reads_flow_s_own_columns(tmp_path):
    """解析器要能从 `Goals.DAT` 里读出 **Flow 自己算的** 平均值/最小值/最大值。

    期望值取自界面「目标图」那一行（已逐位核对）：
        数值=213169.23  平均值=213911.64  最小值=213169.23  最大值=214405.50
        增量=447.65     标准=1065.85

    ⚠️ 这条测试钉住的是一件要紧事：**Flow 的「平均值」在 API 里读不到**
    （`GetValues()`/`GetValues2()` 返回的都是瞬时值序列，实测逐位相同），
    只写在 `Goals.DAT` 这个文本文件里。而训练行必须和用户界面上看到的一致。
    """
    gd = tmp_path / "Goals.DAT"
    gd.mkdir()
    (gd / "SG CV入口静压.txt").write_text(_GOALS_DAT_FIXTURE, encoding="utf-8")
    # 放一个非目标文件，确认会被跳过（Serv*/global_parameters 也会）
    (gd / "Serv Press.txt").write_text(_GOALS_DAT_FIXTURE, encoding="utf-8")

    stats = R.read_goals_from_goals_dat(tmp_path)
    assert set(stats) == {"SG CV入口静压"}, "非目标文件必须被跳过"
    g = stats["SG CV入口静压"]
    assert g["value"] == pytest.approx(213169.234, rel=1e-9)
    assert g["av_value"] == pytest.approx(213911.637, rel=1e-9)
    assert g["window_min"] == pytest.approx(213169.234, rel=1e-9)
    assert g["window_max"] == pytest.approx(214405.497, rel=1e-9)
    assert g["delta"] == pytest.approx(447.64865, rel=1e-9)
    assert g["criteria"] == pytest.approx(1065.84617, rel=1e-9)
    assert g["progress"] == 100.0
    assert g["criteria_percentage"] == 0.5
    assert g["criteria_type"] == 1
    assert g["n_iterations"] == 160
    # 跨度由 Min/Max 推出
    assert g["spread_rel"] == pytest.approx((214405.497 - 213169.234) / 213911.637, rel=1e-9)


def test_goals_dat_parser_returns_empty_when_missing(tmp_path):
    """缺文件时返回空 dict（调用方据此走 API 退路），**不抛异常**。"""
    assert R.read_goals_from_goals_dat(tmp_path) == {}


def test_goal_values_prefers_the_average_over_the_instantaneous():
    """`goal_values` 默认取平均值；`instantaneous` 才取瞬时值。"""
    stats = {"X": {"value": 10.0, "av_value": 11.0}}
    assert R.goal_values(stats)["X"] == 11.0
    assert R.goal_values(stats, "instantaneous")["X"] == 10.0


def test_goals_dat_parser_handles_the_real_file():
    """真实文件的**不变量**检查（不断言具体数值 —— 那些会随重跑变）。

    真正稳的三条：7 个目标都在、关键字段都读得到、判据百分比是个正数。
    """
    if not (HUMAN_DIR / "Goals.DAT").is_dir():
        pytest.skip("根目录没有 Goals.DAT")
    stats = R.read_goals_from_goals_dat(HUMAN_DIR)
    assert set(stats) == set(fg.GOAL_COLUMNS), "7 个目标必须齐全"
    for name, e in stats.items():
        for key in ("value", "av_value", "window_min", "window_max", "delta", "criteria"):
            assert e[key] is not None, f"{name}.{key} 读不到"
        assert e["criteria_percentage"] is not None and e["criteria_percentage"] > 0, name
        assert e["av_value"] == pytest.approx(
            e["av_value"], abs=0), f"{name} 平均值不是有限数"
    # 振荡幅度：`SG 力矩Z` 必须是最大的那个（物理上确定 —— 涡脱落直接打力矩）
    spans = {k: v["spread_rel"] for k, v in stats.items()}
    assert max(spans, key=lambda k: spans[k]) == "SG 力矩Z", spans


def test_the_reference_face_selections_resolve_against_the_real_config():
    """拿真配置跑一遍解析，确认每个目标都指到了**该指的面**。

    期望值来自人工 `1.xmlconfig` 的 Faces_Keys（经面 ID 表反解）：
    密比压 = 大垫片 0/1 + 蝶板 54 + 密封圈 0；蝶板法向压力 = 蝶板 2；力矩Z = 阀轴 2。
    """
    import flow_project as fpj
    want = {
        "SG 蝶板法向压力": [("11蝶板", 2)],
        "SG 密比压 平均": [("08大垫片", 0), ("08大垫片", 1), ("11蝶板", 54), ("09密封圈", 0)],
        "SG 密比压 最大": [("08大垫片", 0), ("08大垫片", 1), ("11蝶板", 54), ("09密封圈", 0)],
        "SG 力矩Z": [("04阀轴", 2)],
    }
    goals = {g["name"]: g for g in fpj.load_reference()["goals"]}
    for name, expected in want.items():
        got = [(e["component"], i)
               for e in goals[name]["face_selection"] for i in e["indices"]]
        assert got == expected, name
        assert goals[name]["expect_faces"] == len(expected)


# ---------------------------------------------------------------- 输入校验

def test_geometry_verified_gate_accepts_both_report_formats():
    """两种报告都要认：ps1 包装器写的带 status，Run-OneDesign 写的没有 status。"""
    R.assert_geometry_verified({"status": "geometry_verified", "completed": True,
                                "physical_mapping_verified": True})
    R.assert_geometry_verified({"completed": True, "physical_mapping_verified": True},
                               Path("_analysis/e2e_x.json"))


def test_geometry_verified_gate_rejects_unverified():
    for bad in ({"status": "failed", "completed": True, "physical_mapping_verified": True},
                {"status": "geometry_verified", "completed": False, "physical_mapping_verified": True},
                {"status": "geometry_verified", "completed": True, "physical_mapping_verified": False},
                {"completed": True}):
        with pytest.raises(RuntimeError, match="几何验收"):
            R.assert_geometry_verified(bad)


# ---------------------------------------------------------------- 报告来源

@pytest.fixture
def fake_mapping_roots(tmp_path, monkeypatch):
    """把 MAPPING_ROOT 指向临时目录，让 _analysis 的定位可测。

    真实布局是 `<SolidWorks-Batch-Parametric-Design-and-CFD-Simulation>/_analysis`（**不在 working 下面**）——
    这里必须显式注入，不能靠数父级。
    """
    monkeypatch.setattr(R, "MAPPING_ROOT", tmp_path)
    return tmp_path


def test_find_mapping_report_prefers_the_run_folder_copy(fake_mapping_roots, tmp_path):
    run = tmp_path / "working" / "seven_variable_trials" / "r1"
    run.mkdir(parents=True)
    (run / "mapping_result.json").write_text('{"status":"geometry_verified"}', encoding="utf-8")
    (tmp_path / "_analysis").mkdir()
    (tmp_path / "_analysis" / "e2e_r1.json").write_text('{"completed":true}', encoding="utf-8")

    path, data = R.find_mapping_report(run)
    assert path == run / "mapping_result.json"
    assert data["status"] == "geometry_verified"


def test_find_mapping_report_falls_back_to_the_e2e_report(fake_mapping_roots, tmp_path):
    """本机跑不了 .ps1，所以 Run-OneDesign.exe 那条链只有 _analysis/e2e_<run>.json。

    ⚠️ `_analysis` 在 SolidWorks-Batch-Parametric-Design-and-CFD-Simulation 根下，**不在 working 下** —— 数父级会数错。
    """
    run = tmp_path / "working" / "seven_variable_trials" / "r2"
    run.mkdir(parents=True)
    (tmp_path / "_analysis").mkdir()
    (tmp_path / "_analysis" / "e2e_r2.json").write_text(
        '{"completed":true,"physical_mapping_verified":true}', encoding="utf-8")

    path, data = R.find_mapping_report(run)
    assert path.name == "e2e_r2.json"
    assert data["physical_mapping_verified"] is True


def test_find_mapping_report_raises_with_both_paths_listed(fake_mapping_roots, tmp_path):
    run = tmp_path / "working" / "seven_variable_trials" / "r3"
    run.mkdir(parents=True)
    with pytest.raises(RuntimeError, match="找不到几何验收报告"):
        R.find_mapping_report(run)


def test_load_design_ignores_the_dotnet_type_name_string(tmp_path):
    """e2e 报告的 `input` 是 `"SevenVariableAdapter+Design"` —— 绝不能当成设计点。"""
    run = tmp_path / "run"
    run.mkdir()
    with pytest.raises(RuntimeError, match="--design"):
        R.load_design(run, {"input": "SevenVariableAdapter+Design"})


def test_load_design_accepts_an_explicit_file(tmp_path):
    design = {"c_mm": 32, "e_mm": 3.7, "phi_deg": 8.25, "alpha_deg": 35.5,
              "Dmax_mm": 191.31786408589767, "bm_mm": 7.5, "ds_mm": 45}
    path = tmp_path / "design.json"
    path.write_text(json.dumps(design), encoding="utf-8")
    run = tmp_path / "run"
    run.mkdir()
    assert R.load_design(run, {}, path) == {k: float(v) for k, v in design.items()}


def test_the_real_baseline_design_file_is_valid():
    """基准设计必须能直接喂进去 —— golden replay 靠它跟 1/1.info.json 逐值比对。"""
    baseline = MAPPING / "config" / "design_baseline_v6.json"
    design = json.loads(baseline.read_text(encoding="utf-8"))
    assert list(design) == list(fg.DESIGN_COLUMNS)
    assert design["c_mm"] == 32.0 and design["phi_deg"] == 8.25
    assert abs(design["Dmax_mm"] - 191.31786408589767) < 1e-9


def test_load_design_prefers_design_input_json(tmp_path):
    run = tmp_path / "run"
    run.mkdir()
    design = {"c_mm": 31, "e_mm": 4, "phi_deg": 8.5, "alpha_deg": 34.5,
              "Dmax_mm": 189, "bm_mm": 7.2, "ds_mm": 44}
    (run / "design_input.json").write_text(json.dumps(design), encoding="utf-8")
    assert R.load_design(run, {"input": {"c_mm": 999}}) == {k: float(v) for k, v in design.items()}


def test_load_design_falls_back_to_mapping_input(tmp_path):
    run = tmp_path / "run"
    run.mkdir()
    design = {"c_mm": 31, "e_mm": 4, "phi_deg": 8.5, "alpha_deg": 34.5,
              "Dmax_mm": 189, "bm_mm": 7.2, "ds_mm": 44}
    assert R.load_design(run, {"input": design}) == {k: float(v) for k, v in design.items()}


def test_load_design_rejects_incomplete_input(tmp_path):
    run = tmp_path / "run"
    run.mkdir()
    with pytest.raises(RuntimeError, match="设计输入缺"):
        R.load_design(run, {"input": {"c_mm": 31}})


# ---------------------------------------------------------------- 状态机

def test_assembly_lookup_ignores_lock_files_and_rejects_two_real_ones(tmp_path):
    """`~$a.SLDASM` 会被 `*.SLDASM` 匹配到 —— 必须过滤掉。

    这正是 C# 那边 `Directory.GetFiles(folder,"*.SLDASM").Single()` 栽过的地方：
    匹配到两个就抛「序列包含一个以上的元素」，表现成适配器内部报错，看不出是文件问题。
    """
    (tmp_path / "a.SLDASM").write_bytes(b"x")
    assert R.assembly_of(tmp_path).name == "a.SLDASM"

    (tmp_path / "~$a.SLDASM").write_bytes(b"\x00" * 12)
    assert R.assembly_of(tmp_path).name == "a.SLDASM", "锁文件不该被算作第二个装配体"

    (tmp_path / "b.SLDASM").write_bytes(b"x")
    with pytest.raises(RuntimeError, match="恰好 1 个装配体"):
        R.assembly_of(tmp_path)


def test_geometry_hashes_ignore_lock_files(tmp_path):
    """`*.SLDPRT` 会匹配到 `~$xxx.SLDPRT` 锁文件 —— 必须排除。

    实测踩过：文档打开时锁文件在、关闭后消失，于是记录下来的哈希表**永远**对不上磁盘，
    每一次 --resume 都被自己的守卫拒掉，而 7 个真实 CAD 文件一个字节都没动。
    """
    (tmp_path / "a.SLDASM").write_bytes(b"asm")
    (tmp_path / "b.SLDPRT").write_bytes(b"part")
    before = R.geometry_hashes(tmp_path)
    assert set(before) == {"a.SLDASM", "b.SLDPRT"}

    (tmp_path / "~$b.SLDPRT").write_bytes(b"locklocklock")
    assert R.geometry_hashes(tmp_path) == before, "锁文件不得进入哈希表"

    state = {"stages": {"cad_rebuild": {"ok": True, "hashes": before}}}
    assert R.resume_is_safe(tmp_path, state)[0], "有锁文件时也不该拒绝续跑"


def test_resume_compares_against_the_latest_mutating_stage(tmp_path):
    """只比对**最后一个**改动几何的阶段。

    `cad_rebuild` 记录之后，`lids` 会往装配体里加封盖并保存 —— 装配体哈希必然变。
    拿 `cad_rebuild` 的旧记录去比，等于每次续跑都必然被自己拒掉。
    """
    (tmp_path / "a.SLDASM").write_bytes(b"before-lids")
    cad = R.hash_many([tmp_path / "a.SLDASM"])

    (tmp_path / "a.SLDASM").write_bytes(b"after-lids")   # Create Lids 保存了装配体
    (tmp_path / "封盖1.SLDPRT").write_bytes(b"cap")
    lids = R.hash_many([tmp_path / "a.SLDASM", tmp_path / "封盖1.SLDPRT"])

    state = {"stages": {"cad_rebuild": {"ok": True, "hashes": cad},
                        "lids": {"ok": True, "hashes": lids}}}
    ok, why = R.resume_is_safe(tmp_path, state)
    assert ok, why
    assert "lids" in why

    # 封盖被动过 -> 必须拒
    (tmp_path / "封盖1.SLDPRT").write_bytes(b"tampered")
    ok, why = R.resume_is_safe(tmp_path, state)
    assert not ok and "封盖1" in why


def test_resume_reports_missing_files(tmp_path):
    (tmp_path / "a.SLDASM").write_bytes(b"x")
    state = {"stages": {"lids": {"ok": True,
                                 "hashes": {"a.SLDASM": "deadbeef", "没了.SLDPRT": "cafe"}}}}
    ok, why = R.resume_is_safe(tmp_path, state)
    assert not ok and "缺失" in why


def test_resume_refuses_when_geometry_moved(tmp_path):
    (tmp_path / "a.SLDASM").write_bytes(b"x")
    hashes = R.hash_many([tmp_path / "a.SLDASM"])
    state = {"stages": {"cad_rebuild": {"ok": True, "hashes": hashes}}}
    assert R.resume_is_safe(tmp_path, state)[0]

    (tmp_path / "a.SLDASM").write_bytes(b"tampered")
    ok, why = R.resume_is_safe(tmp_path, state)
    assert not ok and "不符" in why


def test_resume_is_fine_without_recorded_hashes(tmp_path):
    assert R.resume_is_safe(tmp_path, {"stages": {}})[0]


def test_persist_is_the_hash_source_after_the_pre_solve_save(tmp_path):
    """S9 之后那次「把工程落盘」的保存会改写全部零件/装配体文件。

    它的哈希必须压过 `cap_binding` 的旧记录 —— 否则停在求解前那一次跑完，
    续跑求解时会被守卫拿**保存前**的哈希拒掉，而那一次保存是**正常**的
    （每跑到求解前都会保存一次，为了把 Flow 工程留在装配体里）。
    """
    (tmp_path / "a.SLDASM").write_bytes(b"after-cap-binding")
    cap = R.hash_many([tmp_path / "a.SLDASM"])

    (tmp_path / "a.SLDASM").write_bytes(b"after-persist-save")
    persisted = R.hash_many([tmp_path / "a.SLDASM"])

    state = {"stages": {"cap_binding": {"ok": True, "hashes": cap},
                        "persist": {"ok": True, "hashes": persisted}}}
    ok, why = R.resume_is_safe(tmp_path, state)
    assert ok, why
    assert "persist" in why

    # 只拿 cap_binding 那一份去比必然被拒 —— 证明起作用的确实是 persist
    only_cap = {"stages": {"cap_binding": {"ok": True, "hashes": cap}}}
    assert not R.resume_is_safe(tmp_path, only_cap)[0]


def test_persist_sits_between_cap_binding_and_solve():
    """顺序不能错。`resume_is_safe` 按 MUTATING_STAGES 的顺序取**最后一个**有哈希的阶段：

    * persist 排在 `cap_binding` **后面** → 停在求解前那一次用保存之后的哈希
    * persist 排在 `solve` **前面** → 求解完成后仍是 solve 说了算
      （它的哈希在 `wait_for_result` 之后才记）
    """
    order = list(R.MUTATING_STAGES)
    assert order.index("cap_binding") < order.index("persist") < order.index("solve")


def test_record_and_load_state_round_trip(tmp_path):
    state = R.load_state(tmp_path)
    R._record(state, tmp_path, "gate")
    R._record(state, tmp_path, "open_assembly", {"configuration": "开度45°"})
    again = R.load_state(tmp_path)
    assert again["stages"]["gate"]["ok"] is True
    assert again["stages"]["open_assembly"]["configuration"] == "开度45°"
    assert "updated_utc" in again


# ---------------------------------------------------------------- 物理闸门

def _ref():
    import flow_project as fpj
    return fpj.load_reference()


def test_physics_gate_accepts_the_human_values():
    goals = {"SG CV入口静压": 204866.94744261276, "SG CV出口静压": 101325.0,
             "SG CV入口端面体积流量": 0.10270053171754644}
    assert R.physics_gate(goals, _ref())["ok"]


def test_physics_gate_rejects_a_wrong_outlet_pressure():
    """出口静压就是边界条件，必须精确 —— 差一点说明边界没写对。"""
    goals = {"SG CV入口静压": 204866.947, "SG CV出口静压": 101325.5,
             "SG CV入口端面体积流量": 0.1027}
    result = R.physics_gate(goals, _ref())
    assert not result["ok"] and any("出口静压" in p for p in result["problems"])


def test_physics_gate_rejects_negative_flow_or_pressure_drop():
    flipped = {"SG CV入口静压": 101325.0, "SG CV出口静压": 101325.0,
               "SG CV入口端面体积流量": -0.1}
    result = R.physics_gate(flipped, _ref())
    assert not result["ok"]
    assert any("压差" in p for p in result["problems"])
    assert any("体积流量" in p for p in result["problems"])


# ---------------------------------------------------------------- 训练表

def test_training_row_written_locally_even_if_dataset_is_locked(tmp_path):
    """主表被 Excel 占着时，降级为只写 run 本地副本，**不算失败**。"""
    row = fg.build_training_row(
        {"c_mm": 32, "e_mm": 3.7, "phi_deg": 8.25, "alpha_deg": 35.5,
         "Dmax_mm": 191.3, "bm_mm": 7.5, "ds_mm": 45},
        {"SG CV入口静压": 204866.94744261276, "SG CV出口静压": 101325.0,
         "SG CV入口端面体积流量": 0.10270053171754644},
        998.2)
    result = R.write_training_row(row, tmp_path, tmp_path / "nope" / "deep" / "x.xlsx")
    local = tmp_path / "training_sample.csv"
    assert local.is_file()
    header = local.read_text(encoding="utf-8-sig").splitlines()[0].split(",")
    assert header == list(fg.TRAINING_COLUMNS)
    assert "run_csv" in result


def test_training_dataset_deduplicates_an_identical_design_and_result(tmp_path):
    row = fg.build_training_row(
        {"c_mm": 32, "e_mm": 3.7, "phi_deg": 8.25, "alpha_deg": 35.5,
         "Dmax_mm": 191.3, "bm_mm": 7.5, "ds_mm": 45},
        {"SG CV入口静压": 204866.94744261276, "SG CV出口静压": 101325.0,
         "SG CV入口端面体积流量": 0.10270053171754644},
        998.2)
    path = tmp_path / "training.xlsx"

    first = R._append_xlsx(path, row)
    second = R._append_xlsx(path, row)

    assert first == {"appended": True, "deduplicated": False}
    assert second["appended"] is False and second["deduplicated"] is True
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.active.max_row == 2
    finally:
        workbook.close()


def test_training_dataset_rejects_same_design_with_different_targets(tmp_path):
    row = fg.build_training_row(
        {"c_mm": 32, "e_mm": 3.7, "phi_deg": 8.25, "alpha_deg": 35.5,
         "Dmax_mm": 191.3, "bm_mm": 7.5, "ds_mm": 45},
        {"SG CV入口静压": 204866.94744261276, "SG CV出口静压": 101325.0,
         "SG CV入口端面体积流量": 0.10270053171754644},
        998.2)
    path = tmp_path / "training.xlsx"
    R._append_xlsx(path, row)
    conflicting = dict(row)
    conflicting[fg.GOAL_COLUMNS[0]] += 1.0    # 篡改第一个目标 —— 不是设计输入

    with pytest.raises(RuntimeError, match="相同七变量"):
        R._append_xlsx(path, conflicting)


def test_training_dataset_ignores_the_sample_id_when_deduplicating(tmp_path):
    """**样本序号是元数据，不参与「是不是同一行」的判定。**

    回归：把编号算进整行比较的话，一条老行（编号为空）和一条新行（编号=7）
    即便七变量与目标值逐位相同，也会被判成「同输入不同目标」而拒写 ——
    表现是整批数据一行都进不去，而脚本每步都报成功。
    """
    design = {"c_mm": 32, "e_mm": 3.7, "phi_deg": 8.25, "alpha_deg": 35.5,
              "Dmax_mm": 191.3, "bm_mm": 7.5, "ds_mm": 45}
    goals = {"SG CV入口静压": 204866.94744261276, "SG CV出口静压": 101325.0,
             "SG CV入口端面体积流量": 0.10270053171754644}
    legacy = fg.build_training_row(design, goals, 998.2)          # 编号留空（老行）
    batched = fg.build_training_row(design, goals, 998.2, sample_id=7)
    path = tmp_path / "training.xlsx"

    assert R._append_xlsx(path, legacy) == {"appended": True, "deduplicated": False}
    result = R._append_xlsx(path, batched)

    assert result["appended"] is False and result["deduplicated"] is True, \
        "只剩编号不同 → 仍应判重复，不得报「同输入不同目标」"


# ---------------------------------------------------------------- 模块卫生

def test_module_imports_without_com():
    source = (MAPPING / "scripts" / "Run-FlowSample.py").read_text(encoding="utf-8")
    header = source.split("def now", 1)[0]
    assert "import win32com" not in header
    assert "import pythoncom" not in header


def test_no_dead_if_false_leftovers():
    """写这段代码时反复返工，`if False else` 这种残留必须清干净。"""
    for name in ("Run-FlowSample.py", "flow_geometry.py", "flow_project.py",
                 "flow_session.py", "sw_api.py"):
        text = (MAPPING / "scripts" / name).read_text(encoding="utf-8")
        assert "if False" not in text, f"{name} 里有 if False 残留"


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
