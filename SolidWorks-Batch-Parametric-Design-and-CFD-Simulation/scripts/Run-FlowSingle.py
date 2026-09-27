#!/usr/bin/env python3
"""Preflight and solve one geometry-verified parameterized Flow Simulation run."""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import pythoncom
import win32com.client as win32
from win32com.client import gencache

PROJECT_ROOT = Path(__file__).resolve().parents[2]
MAPPING_ROOT = PROJECT_ROOT / "SolidWorks-Batch-Parametric-Design-and-CFD-Simulation"
RUNS_ROOT = (MAPPING_ROOT / "working" / "seven_variable_trials").resolve()
sys.path.insert(0, str(PROJECT_ROOT))
import flow_transfer as ft  # noqa: E402
import machine_paths as machine_paths  # noqa: E402

PROJECT_NAME = "流体力学仿真"
CONFIGURATION_NAME = "开度45°"
AUTO_DISMISS_MESSAGES = 564
ASSEMBLY_MESSAGE_SECONDS = 519
EXPECTED = {
    "入口速度 2": 2,
    "出口静压 2": 2,
    "SG CV入口静压": 12,
    "SG CV出口静压": 12,
    "SG CV入口端面体积流量": 12,
    "SG 蝶板法向压力": 12,
    "SG 密比压 平均": 12,
    "SG 密比压 最大": 12,
    "SG 力矩Z": 12,
}
TRAINING_HEADERS = [
    "c_mm", "e_mm", "phi_deg", "alpha_deg", "Dmax_mm", "bm_mm", "ds_mm",
    "Cv", "SG 力矩Z", "SG 蝶板法向压力", "SG 密比压 平均", "SG 密比压 最大",
]


def now():
    return datetime.now(timezone.utc).isoformat()


def inside(path: Path, parent: Path) -> Path:
    path = path.resolve()
    if path == parent or parent not in path.parents:
        raise RuntimeError(f"Path must be below {parent}: {path}")
    return path


def get_attr(obj, name):
    value = getattr(obj, name)
    return value() if callable(value) else value


def inspect_project(project):
    features = ft.read_features(project)
    found = {str(item["name"]): int(item["type"]) for item in features}
    missing = [name for name, kind in EXPECTED.items() if found.get(name) != kind]
    reference_errors = []
    for item in features:
        if item["name"] in EXPECTED and not item.get("references"):
            reference_errors.append(item["name"])
    files = project.ProjectFiles
    paths = {}
    for name in ("ProjectDirectory", "GEOMFile", "CPTFile", "FLDFile"):
        try:
            paths[name] = str(get_attr(files, name))
        except Exception as exc:
            paths[name] = None
            paths[name + "_error"] = str(exc)
    return features, missing, reference_errors, paths


def write_report(path, report):
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def rebuild_features(project):
    results = []
    enum = project.EnumFeatures()
    enum.Reset()
    while True:
        feature = enum.Next()
        if feature is None:
            break
        ok = bool(feature.Rebuild(project))
        item = {"name": str(feature.Name), "type": int(feature.Type), "ok": ok}
        if not ok:
            try:
                item["error"] = str(project.GetLastRebuildError())
            except Exception as exc:
                item["error"] = f"unavailable: {exc}"
        results.append(item)
    return results


def feature_objects(project):
    items = {}
    enum = project.EnumFeatures()
    enum.Reset()
    while True:
        feature = enum.Next()
        if feature is None:
            return items
        items[str(feature.Name)] = feature


def name_inner_cap_face(sw, assembly, cap_prefix, other_prefix, stable_name):
    sw_types = gencache.EnsureModule("{83A33D31-27C5-11CE-BFD4-00400513BB57}", 0, 34, 0)
    components = list(assembly.GetComponents(False) or [])
    cap = next((c for c in components if str(c.Name2).startswith(cap_prefix + "-")), None)
    other = next((c for c in components if str(c.Name2).startswith(other_prefix + "-")), None)
    if cap is None or other is None:
        raise RuntimeError(f"Could not find native cap components {cap_prefix}/{other_prefix}")
    def transform_point(point, transform):
        data = list(transform.ArrayData)
        scale = data[12] if len(data) > 12 else 1.0
        x, y, z = point
        return [scale * (x * data[0] + y * data[3] + z * data[6]) + data[9],
                scale * (x * data[1] + y * data[4] + z * data[7]) + data[10],
                scale * (x * data[2] + y * data[5] + z * data[8]) + data[11]]
    other_origin = transform_point([0.0, 0.0, 0.0], other.Transform2)
    part_path = Path(str(ft.safe_get(assembly, "GetPathName", "PathName"))).parent / f"{cap_prefix}.SLDPRT"
    part = sw.GetOpenDocumentByName(str(part_path))
    if part is None:
        raise RuntimeError(f"Could not access part document for {cap.Name2}")
    candidates = []
    for body in list(part.GetBodies2(0, False) or []):
        for raw_face in list(body.GetFaces() or []):
            face = sw_types.IFace(raw_face._oleobj_)
            surface = get_attr(face, "GetSurface")
            if not bool(get_attr(surface, "IsPlane")):
                continue
            box = list(get_attr(face, "GetBox"))
            local_center = [(box[0] + box[3]) / 2, (box[1] + box[4]) / 2, (box[2] + box[5]) / 2]
            world = transform_point(local_center, cap.Transform2)
            distance2 = sum((world[i] - other_origin[i]) ** 2 for i in range(3))
            candidates.append((distance2, -float(get_attr(face, "GetArea")), face, local_center, world))
    if not candidates:
        raise RuntimeError(f"No planar face found on {cap.Name2}")
    candidates.sort(key=lambda item: (item[0], item[1]))
    _, _, chosen, local_center, world = candidates[0]
    existing = str(part.GetEntityName(chosen) or "")
    if existing != stable_name:
        if not bool(part.SetEntityName(chosen, stable_name)):
            raise RuntimeError(f"Could not name inner face {stable_name} on {cap.Name2}")
        save_result = bool(get_attr(part, "Save"))
    else:
        save_result = None
    audit = {"component": str(cap.Name2), "face_name": stable_name,
            "previous_name": existing, "local_center_m": local_center,
            "world_center_m": world, "planar_candidate_count": len(candidates),
            "part_save_result": save_result}
    return audit, cap


def select_named_cap_face(model, component, face_name, null_dispatch):
    title = str(ft.safe_get(model, "GetTitle", "Title") or "")
    names = [f"{face_name}@{component.Name2}@{title}",
             f"{face_name}@{component.Name2}", face_name]
    for selection_name in names:
        model.ClearSelection2(True)
        if bool(model.Extension.SelectByID2(
                selection_name, "FACE", 0, 0, 0, False, 0, null_dispatch, 0)):
            return selection_name
    return None


def repair_endpoint_references(project, sw):
    """Rebind only invalid inlet/outlet topology; preserve feature names and parameters."""
    groups = [
        ("封盖1<1>", "封盖1", "封盖2", "FLOW_INLET_INNER",
         ["入口速度 2", "SG CV入口静压", "SG CV入口端面体积流量"]),
        ("封盖2<1>", "封盖2", "封盖1", "FLOW_OUTLET_INNER",
         ["出口静压 2", "SG CV出口静压"]),
    ]
    features = feature_objects(project)
    native_model = sw.ActiveDoc
    if native_model is None:
        raise RuntimeError("SolidWorks native API returned no active document")
    null_dispatch = win32.VARIANT(pythoncom.VT_DISPATCH, None)
    audit = []
    for component, native_prefix, other_prefix, stable_name, names in groups:
        boundary = features[names[0]]
        if bool(boundary.Rebuild(project)):
            chosen = list(boundary.GetInterface("ITopologyBasedFeature").GetReferencesNames())
            audit.append({"component": component, "changed": False, "reference": chosen})
            continue
        named, native_component = name_inner_cap_face(
            sw, native_model, native_prefix, other_prefix, stable_name)
        chosen_face = stable_name
        attempts = []
        topo = boundary.GetInterface("ITopologyBasedFeature")
        topo.RemoveAllReferencies()
        native_model.ClearSelection2(True)
        selected_name = select_named_cap_face(native_model, native_component, chosen_face, null_dispatch)
        selected = selected_name is not None
        added = bool(topo.UpdateReferenciesFromSelection()) if selected else False
        rebuilt = bool(boundary.Rebuild(project)) if added else False
        attempts.append({"face": chosen_face, "selected": selected, "added": added, "rebuilt": rebuilt})
        if not rebuilt:
            raise RuntimeError(f"Geometrically identified inner face is not a Flow boundary: {named}")
        updated = []
        for name in names[1:]:
            feature = features[name]
            goal_topo = feature.GetInterface("ITopologyBasedFeature")
            goal_topo.RemoveAllReferencies()
            native_model.ClearSelection2(True)
            selected_name = select_named_cap_face(native_model, native_component, chosen_face, null_dispatch)
            selected = selected_name is not None
            added = bool(goal_topo.UpdateReferenciesFromSelection()) if selected else False
            rebuilt = bool(feature.Rebuild(project)) if added else False
            updated.append({"name": name, "selected": selected, "added": added, "rebuilt": rebuilt})
            if not rebuilt:
                raise RuntimeError(f"Could not bind {name} to {chosen_face}@{component}")
        audit.append({"component": component, "changed": True,
                      "reference": f"{chosen_face}@{component}", "attempts": attempts,
                      "named_face": named, "dependent_goals": updated})
    return audit


def read_problem_flags(project_dir):
    configs = sorted(project_dir.glob("*.xmlconfig"))
    if not configs:
        raise RuntimeError(f"No .xmlconfig generated in {project_dir}")
    text = configs[0].read_text(encoding="utf-8", errors="replace")
    def flag(name):
        match = re.search(rf'<{name}\s+value="([0-9]+)"', text)
        return int(match.group(1)) if match else None
    return {"path": str(configs[0]), "ProblemType": flag("ProblemType"),
            "FlowSpaceType": flag("FlowSpaceType")}


def solver_running():
    result = subprocess.run(["tasklist", "/FI", "IMAGENAME eq EFDsolver.exe"],
                            capture_output=True, text=True, errors="replace")
    return "efdsolver.exe" in result.stdout.lower()


def wait_for_result(fld_path, started_epoch, timeout_seconds, report_path, report):
    deadline = time.time() + timeout_seconds
    stable = 0
    last = None
    while time.time() < deadline:
        running = solver_running()
        stat = fld_path.stat() if fld_path.is_file() else None
        fresh = bool(stat and stat.st_mtime >= started_epoch - 2 and stat.st_size > 0)
        signature = (stat.st_size, stat.st_mtime_ns) if stat else None
        stable = stable + 1 if signature is not None and signature == last else 0
        last = signature
        report["solver_monitor"] = {
            "efdsolver_running": running, "fld_exists": stat is not None,
            "fld_fresh": fresh, "fld_size": stat.st_size if stat else None,
            "checked_utc": now(),
        }
        write_report(report_path, report)
        if fresh and not running and stable >= 2:
            return
        time.sleep(5)
    raise TimeoutError(f"Flow solve did not produce a completed fresh FLD within {timeout_seconds} seconds")


def read_goals(nca, fld_path):
    handler = nca.LoadFDAResultFile(str(fld_path), True)
    if handler is None:
        raise RuntimeError("LoadFDAResultFile returned no result handler")
    results = handler.GetGoalsCalculationResults2()
    if results is None:
        raise RuntimeError("GetGoalsCalculationResults2 returned no results")
    enum = results.GetGoalsEnum()
    enum.Reset()
    goals = []
    for _ in range(1000):
        goal = enum.Next()
        if goal is None:
            break
        value = float(goal.GetLastCalculatedValue())
        history = list(goal.GetValues2() or [])
        goals.append({"name": str(goal.GetGoalName()), "last_value_si": value,
                      "finite": math.isfinite(value), "history_si": history,
                      "history_count": len(history)})
    else:
        raise RuntimeError("Goal enumeration exceeded safety limit")
    return goals


def append_training_workbook(path, run_key, training_row):
    """Atomically append one numeric training row; a hidden run index makes retries idempotent."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file():
        workbook = load_workbook(path)
        if "训练数据" not in workbook.sheetnames or "_runs" not in workbook.sheetnames:
            raise RuntimeError(f"Training workbook schema is invalid: {path}")
        sheet = workbook["训练数据"]
        index = workbook["_runs"]
        existing_headers = [sheet.cell(1, column).value for column in range(1, len(TRAINING_HEADERS) + 1)]
        if existing_headers != TRAINING_HEADERS:
            raise RuntimeError(f"Training workbook headers differ from the audited schema: {path}")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "训练数据"
        sheet.append(TRAINING_HEADERS)
        index = workbook.create_sheet("_runs")
        index.append(["run_key", "data_row"])
        index.sheet_state = "hidden"

    for row_number in range(2, index.max_row + 1):
        if index.cell(row_number, 1).value == run_key:
            data_row = int(index.cell(row_number, 2).value)
            existing = [sheet.cell(data_row, column).value for column in range(1, len(TRAINING_HEADERS) + 1)]
            desired = [training_row[name] for name in TRAINING_HEADERS]
            if existing != desired:
                raise RuntimeError(f"Run key already exists with different values in {path}: {run_key}")
            workbook.close()
            return data_row, False

    values = [training_row[name] for name in TRAINING_HEADERS]
    if any(value is None or not math.isfinite(float(value)) for value in values):
        raise RuntimeError(f"Training row contains missing or non-finite values: {training_row}")
    sheet.append(values)
    data_row = sheet.max_row
    index.append([run_key, data_row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{data_row}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [12, 12, 12, 12, 14, 12, 12, 14, 16, 20, 18, 18]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=len(TRAINING_HEADERS)):
        for cell in row:
            cell.number_format = "0.000000"
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, path)
    return data_row, True


def export_results(run, mapping, goals, density, dataset_xlsx):
    by_name = {g["name"]: g["last_value_si"] for g in goals}
    required = ["SG CV入口静压", "SG CV出口静压", "SG CV入口端面体积流量"]
    missing = [name for name in required if name not in by_name]
    if missing:
        raise RuntimeError(f"Required CV goals are absent from result: {missing}")
    pin, pout = by_name[required[0]], by_name[required[1]]
    q = abs(by_name[required[2]])
    delta_p = pin - pout
    if not all(math.isfinite(x) for x in (pin, pout, q, delta_p)) or delta_p <= 0:
        raise RuntimeError(f"CV inputs invalid: Pin={pin}, Pout={pout}, Q={q}")
    cv = (q * 60.0 / 0.003785411784) * math.sqrt((density / 1000.0) / (delta_p / 6894.757293168))
    compact = {
        "schema_version": 1, "generated_utc": now(), "run_folder": str(run),
        "design_variables": mapping.get("input", {}), "opening_deg": 45.0,
        "density_kg_m3": density, "volume_flow_m3_s": q,
        "inlet_static_pressure_pa": pin, "outlet_static_pressure_pa": pout,
        "delta_pressure_pa": delta_p, "Cv": cv, "goals": goals,
        "units": "SI except Cv (US customary valve flow coefficient)",
    }
    json_path = run / "simulation_results.json"
    csv_path = run / "simulation_results.csv"
    training_path = run / "training_sample.csv"
    write_report(json_path, compact)
    row = dict(mapping.get("input", {}))
    row.update({"opening_deg": 45.0, "density_kg_m3": density,
                "volume_flow_m3_s": q, "inlet_static_pressure_pa": pin,
                "outlet_static_pressure_pa": pout, "delta_pressure_pa": delta_p, "Cv": cv})
    for goal in goals:
        row[f"goal::{goal['name']}"] = goal["last_value_si"]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(row))
        writer.writeheader()
        writer.writerow(row)
    variable_names = TRAINING_HEADERS[:7]
    objective_names = TRAINING_HEADERS[7:]
    training_row = {name: mapping.get("input", {}).get(name) for name in variable_names}
    training_row["Cv"] = cv
    for name in objective_names[1:]:
        training_row[name] = by_name.get(name)
    with training_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=variable_names + objective_names)
        writer.writeheader()
        writer.writerow(training_row)
    sample_xlsx = run / "training_sample.xlsx"
    append_training_workbook(sample_xlsx, str(run.resolve()), training_row)
    dataset_row, dataset_appended = append_training_workbook(dataset_xlsx, str(run.resolve()), training_row)
    return compact, json_path, csv_path, training_path, sample_xlsx, dataset_xlsx, dataset_row, dataset_appended


def connect_or_start_flow():
    api = ft.connect()
    interactive = ft.safe_get(api, "Attach2RunningObject")
    if interactive is not None:
        return api, interactive, False
    try:
        solidworks_exe = machine_paths.solidworks_exe()
        bincfw = machine_paths.flow_bincfw()
    except RuntimeError:
        api.UnloadProductAPI()
        raise
    interactive = api.RunProduct2(str(solidworks_exe), str(bincfw))
    if interactive is None:
        api.UnloadProductAPI()
        raise RuntimeError("Flow API RunProduct2 returned no InteractiveApplication")
    return api, interactive, True


def configure_temporary_font_substitution():
    sw = None
    for attempt in range(60):
        try:
            sw = win32.GetActiveObject("SldWorks.Application")
        except Exception:
            sw = None
        if sw is not None:
            break
        if attempt < 59:
            time.sleep(1)
    if sw is None:
        raise RuntimeError("Could not access SolidWorks preferences before opening documents")
    previous = {
        "auto_dismiss": bool(sw.GetUserPreferenceToggle(AUTO_DISMISS_MESSAGES)),
        "seconds": int(sw.GetUserPreferenceIntegerValue(ASSEMBLY_MESSAGE_SECONDS)),
    }
    set_toggle = bool(sw.SetUserPreferenceToggle(AUTO_DISMISS_MESSAGES, True))
    set_seconds = bool(sw.SetUserPreferenceIntegerValue(ASSEMBLY_MESSAGE_SECONDS, 1))
    current = {
        "auto_dismiss": bool(sw.GetUserPreferenceToggle(AUTO_DISMISS_MESSAGES)),
        "seconds": int(sw.GetUserPreferenceIntegerValue(ASSEMBLY_MESSAGE_SECONDS)),
    }
    if not current["auto_dismiss"] or current["seconds"] != 1:
        raise RuntimeError("SolidWorks did not accept the automatic message-dismiss preferences")
    return sw, previous, {"toggle_set": set_toggle, "seconds_set": set_seconds, "current": current}


def restore_preferences(sw, previous):
    if sw is None or previous is None:
        return None
    sw.SetUserPreferenceIntegerValue(ASSEMBLY_MESSAGE_SECONDS, int(previous["seconds"]))
    sw.SetUserPreferenceToggle(AUTO_DISMISS_MESSAGES, bool(previous["auto_dismiss"]))
    return {
        "auto_dismiss": bool(sw.GetUserPreferenceToggle(AUTO_DISMISS_MESSAGES)),
        "seconds": int(sw.GetUserPreferenceIntegerValue(ASSEMBLY_MESSAGE_SECONDS)),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("run_folder")
    parser.add_argument("--solve", action="store_true")
    parser.add_argument("--validate-flow-rebuild", action="store_true",
                        help="Diagnostic only: update geometry and rebuild all Flow features, but do not mesh/solve")
    parser.add_argument("--allow-unparameterized-diagnostic", action="store_true")
    parser.add_argument("--timeout-min", type=float, default=120.0)
    parser.add_argument("--rho-kg-m3", type=float, default=1000.0)
    parser.add_argument("--dataset-xlsx", default=str(MAPPING_ROOT / "outputs" / "training_dataset.xlsx"))
    args = parser.parse_args()
    run = inside(Path(args.run_folder), RUNS_ROOT)
    dataset_xlsx = Path(args.dataset_xlsx).resolve()
    if dataset_xlsx == MAPPING_ROOT.resolve() or MAPPING_ROOT.resolve() not in dataset_xlsx.parents:
        raise RuntimeError(f"Dataset workbook must stay inside {MAPPING_ROOT}: {dataset_xlsx}")
    if dataset_xlsx.suffix.lower() != ".xlsx":
        raise RuntimeError("Dataset workbook must use the .xlsx extension")
    if args.solve and args.validate_flow_rebuild:
        raise RuntimeError("Use --solve or --validate-flow-rebuild, not both")
    mapping_path = run / "mapping_result.json"
    if mapping_path.is_file():
        mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    elif args.allow_unparameterized_diagnostic and args.validate_flow_rebuild:
        mapping = {"status": "unparameterized_diagnostic", "input": {}}
    else:
        raise RuntimeError("Missing mapping_result.json")
    geometry_verified = (mapping.get("status") == "geometry_verified" and mapping.get("completed") is True
                         and mapping.get("physical_mapping_verified") is True)
    if not geometry_verified and not (args.allow_unparameterized_diagnostic and args.validate_flow_rebuild):
        raise RuntimeError("Flow run requires a completed geometry_verified mapping_result.json")
    assemblies = list(run.glob("*.SLDASM"))
    parts = list(run.glob("*.SLDPRT"))
    expected_parts = {
        "8“D94R3Y-CL600C-03阀体.SLDPRT", "8“D94R3Y-CL600C-04阀轴.SLDPRT",
        "8“D94R3Y-CL600C-08大垫片.SLDPRT", "8“D94R3Y-CL600C-09密封圈.SLDPRT",
        "8“D94R3Y-CL600C-10压板.SLDPRT", "8“D94R3Y-CL600C-11蝶板.SLDPRT",
        "封盖1.SLDPRT", "封盖2.SLDPRT", "封盖3.SLDPRT", "封盖4.SLDPRT",
    }
    actual_parts = {p.name for p in parts}
    if len(assemblies) != 1 or actual_parts != expected_parts:
        raise RuntimeError(f"Expected the v6 assembly and ten exact external parts; got {sorted(actual_parts)}")

    report = {
        "schema_version": 1,
        "mode": "solve" if args.solve else ("flow_rebuild_diagnostic" if args.validate_flow_rebuild else "preflight"),
        "started_utc": now(),
        "run_folder": str(run),
        "assembly": str(assemblies[0]),
        "geometry_report": str(mapping_path),
        "geometry_verified": geometry_verified,
        "project_name_expected": PROJECT_NAME,
        "configuration_expected": CONFIGURATION_NAME,
        "mesh_started": False,
        "solver_started": False,
        "status": "starting",
    }
    output = run / ("flow_run.json" if args.solve else
                    ("flow_rebuild_diagnostic.json" if args.validate_flow_rebuild else "flow_preflight.json"))
    pythoncom.CoInitialize()
    sw = model = nca = interactive = None
    previous_preferences = None
    started_solidworks = False
    try:
        nca, interactive, started_solidworks = connect_or_start_flow()
        report["solidworks_started_by_script"] = started_solidworks
        sw, previous_preferences, preference_set = configure_temporary_font_substitution()
        report["font_dialog_policy"] = {
            "mode": "temporary_replace_all_via_automatic_dismiss",
            "previous": previous_preferences,
            "set": preference_set,
        }
        model = interactive.OpenDocument(str(assemblies[0]), CONFIGURATION_NAME)
        if model is None:
            raise RuntimeError("SolidWorks could not open the parameterized assembly")
        opened = str(ft.safe_get(model, "GetPathName", "PathName") or "")
        if Path(opened).resolve() != assemblies[0].resolve():
            raise RuntimeError(f"SolidWorks opened a different assembly: {opened}")
        document = interactive.ActiveDocument
        configuration = document.ActiveConfiguration
        config_name = str(ft.safe_get(configuration, "Name"))
        report["configuration"] = config_name
        if config_name != CONFIGURATION_NAME:
            raise RuntimeError(f"Unexpected active configuration: {config_name}")
        names = ft.project_names(configuration)
        report["projects"] = names
        if PROJECT_NAME not in names:
            raise RuntimeError(f"Required Flow project is missing: {PROJECT_NAME}")
        project = configuration.ActivateProject(PROJECT_NAME, False)
        if project is None:
            raise RuntimeError("Could not activate the existing Flow project")
        features, missing, reference_errors, project_files = inspect_project(project)
        report["features"] = features
        report["missing_or_wrong_type"] = missing
        report["empty_topology_references"] = reference_errors
        report["project_files"] = project_files
        project_dir_text = project_files.get("ProjectDirectory")
        if not project_dir_text:
            raise RuntimeError("Flow project directory is unavailable")
        project_dir = Path(project_dir_text).resolve()
        report["project_directory_inside_run"] = project_dir == run or run in project_dir.parents
        if not report["project_directory_inside_run"]:
            raise RuntimeError(f"Flow project points outside the run folder: {project_dir}")
        if missing or reference_errors:
            raise RuntimeError("Flow project features or topology references failed validation")
        if args.solve or args.validate_flow_rebuild:
            report["endpoint_reference_policy"] = {
                "mode": "user_verified_fixed_references",
                "inlet": "面<1>@封盖1<1>", "outlet": "面<1>@封盖2<1>",
                "changed": False,
            }
            report["stage"] = "update_flow_geometry_and_configuration"
            native_doc = sw.ActiveDoc
            if native_doc is None:
                raise RuntimeError("Native SolidWorks document unavailable for rebuild")
            report["document_rebuild_result"] = bool(native_doc.ForceRebuild3(False))
            report["update_config_and_data_files_result"] = bool(project.UpdateConfigAndDataFiles())
            report["document_save_result"] = bool(document.Save())
            project_dir = Path(str(get_attr(project.ProjectFiles, "ProjectDirectory"))).resolve()
            flags = read_problem_flags(project_dir)
            report["problem_flags"] = flags
            if flags["ProblemType"] != 1:
                raise RuntimeError(f"Refusing solve: expected internal-flow ProblemType=1, got {flags['ProblemType']}")
            report["stage"] = "rebuild_flow_features_after_geometry_update"
            report["feature_rebuild"] = rebuild_features(project)
            failed_rebuilds = [x for x in report["feature_rebuild"] if not x["ok"]]
            if failed_rebuilds:
                raise RuntimeError(f"Flow feature rebuild failed after geometry update: {failed_rebuilds}")
            if args.validate_flow_rebuild:
                report["status"] = "flow_rebuild_verified"
                report["training_ready"] = False
                return
            fld_path = Path(str(get_attr(project.ProjectFiles, "FLDFile"))).resolve()
            if not (fld_path == run or run in fld_path.parents):
                raise RuntimeError(f"FLD output points outside run folder: {fld_path}")
            report["stage"] = "solve"
            report["mesh_started"] = True
            report["solver_started"] = True
            report["solve_started_utc"] = now()
            write_report(output, report)
            solve_epoch = time.time()
            report["solve2_return"] = bool(project.Solve2(True, True, True, False))
            wait_for_result(fld_path, solve_epoch, args.timeout_min * 60.0, output, report)
            report["stage"] = "read_results"
            report["result_fld"] = str(fld_path)
            goals = read_goals(nca, fld_path)
            report["goals"] = goals
            (compact, json_path, csv_path, training_path, sample_xlsx, dataset_xlsx,
             dataset_row, dataset_appended) = export_results(
                run, mapping, goals, args.rho_kg_m3, dataset_xlsx)
            report["calculated"] = {k: compact[k] for k in (
                "volume_flow_m3_s", "inlet_static_pressure_pa", "outlet_static_pressure_pa",
                "delta_pressure_pa", "density_kg_m3", "Cv")}
            report["exports"] = {"json": str(json_path), "csv": str(csv_path),
                                 "training_sample_csv": str(training_path),
                                 "training_sample_xlsx": str(sample_xlsx),
                                 "training_dataset_xlsx": str(dataset_xlsx),
                                 "training_dataset_row": dataset_row,
                                 "training_dataset_appended": dataset_appended}
            report["training_ready"] = True
            report["status"] = "solve_and_export_completed"
        else:
            report["status"] = "preflight_verified"
    except Exception as exc:
        report["status"] = "preflight_failed" if not args.solve else "solve_failed"
        report["error"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        report["finished_utc"] = now()
        try:
            restored = restore_preferences(sw, previous_preferences)
            if "font_dialog_policy" in report:
                report["font_dialog_policy"]["restored"] = restored
        except Exception as exc:
            report["font_dialog_preference_restore_error"] = str(exc)
        if interactive is not None and model is not None:
            try:
                interactive.CloseActiveDoc()
            except Exception:
                pass
        if started_solidworks and interactive is not None:
            try:
                interactive.ExitApp2()
            except Exception:
                pass
        if nca is not None:
            try:
                nca.UnloadProductAPI()
            except Exception:
                pass
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        pythoncom.CoUninitialize()
    print(json.dumps({"status": report["status"], "report": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
